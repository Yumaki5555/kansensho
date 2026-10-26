"""まとめたデータからExcelを作る。
使い方: python -I build_xlsx.py まとめたデータのフォルダ 出力xlsxのパス
"""
import json
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
data_dir, out_path = sys.argv[1], sys.argv[2]
load = lambda n: json.load(open(os.path.join(data_dir, f"{n}.json"), encoding="utf-8"))  # noqa: E731
annual, ytd, breakdown, notes, sources, world, kouhi = (load(n) for n in ["annual", "ytd2026", "breakdown", "notes", "sources", "world", "kouhi"])
countries, entry, costs, info, others, guide = (load(n) for n in ["countries", "entry", "costs", "info", "others", "guide"])
WORLD = "HIV/エイズ（世界）"
KOUHI = "公費の制度"

FONT = "游ゴシック"
GREEN = "28765B"
f_body = Font(name=FONT, size=10.5)
f_bold = Font(name=FONT, size=10.5, bold=True)
f_title = Font(name=FONT, size=15, bold=True, color=GREEN)
f_sec = Font(name=FONT, size=12, bold=True, color=GREEN)
f_head = Font(name=FONT, size=10.5, bold=True, color="FFFFFF")
f_note = Font(name=FONT, size=9.5, color="5F6B66")
f_link = Font(name=FONT, size=10.5, color="1F5FA8", underline="single")
fill_head = PatternFill("solid", fgColor=GREEN)
fill_soku = PatternFill("solid", fgColor="FFF4D1")
fill_soft = PatternFill("solid", fgColor="EEF5F1")
fill_yes = PatternFill("solid", fgColor="DDEFE4")
fill_no = PatternFill("solid", fgColor="EEEEEE")
fill_cond = PatternFill("solid", fgColor="FBF0D5")
thin = Side(style="thin", color="C9D3CE")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

SHEETS = [("HIV/エイズ", "HIV・エイズ", "HIV/エイズ", "合計", "合計"),
          ("梅毒", "梅毒", "梅毒", "合計", "合計"),
          ("はしか", "はしか", "はしか（麻しん）", "届出数", "届出数"),
          ("結核", "結核", "結核", "新登録患者数", "届出数（発病していない感染者も含む）")]
WHAT = {"HIV/エイズ": "新たに報告されたHIV感染者とエイズ患者の合計（件）",
        "梅毒": "医師が届け出た患者の数（人）",
        "はしか": "医師が届け出た患者の数（人）",
        "結核": "新たに登録された結核患者の数（人）"}
WIDTH = 9  # 文章を横に広げる列数


def numfmt(col, v):
    if isinstance(v, float) or "率" in col or "(%)" in col or "人口10万対" in col:
        return "0.0" if not ("人口10万対" in col and isinstance(v, float) and abs(v * 1000 - round(v * 1000)) < 1e-6 and round(v, 1) != v) else "0.000"
    return "#,##0"


def put(ws, r, c, v, font=f_body, fill=None, fmt=None, align=None, border=True):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = align
    if border:
        cell.border = box
    return cell


def para(ws, r, text, font=f_body, cols=WIDTH, per_line=62):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=cols)
    c = ws.cell(row=r, column=1, value=text)
    c.font = font
    c.alignment = wrap
    ws.row_dimensions[r].height = max(18, 16 * math.ceil(len(text) / per_line))
    return r + 1


def header(ws, r, labels):
    for i, h in enumerate(labels, 1):
        put(ws, r, i, h, f_head, fill_head, align=center)
    ws.row_dimensions[r].height = 30
    return r + 1


def section(ws, r, text):
    ws.cell(row=r, column=1, value=text).font = f_sec
    ws.row_dimensions[r].height = 22
    return r + 1


def notes_of(name, kind, table=None):
    return [n for n in notes if n["病気"] == name and n["種類"] == kind and n["表"] == table]


wb = Workbook()
intro = wb.active
intro.title = "はじめに"
links = {}  # 病気 -> まとめ表から参照するセル

for name, sheet, display, main, ymain in SHEETS:
    ws = wb.create_sheet(sheet)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 40
    for i in range(2, 12):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.cell(row=1, column=1, value=display).font = f_title
    ws.row_dimensions[1].height = 26
    r = para(ws, 2, "数え方：" + notes_of(name, "数え方")[0]["文"], f_note)
    r += 1

    # 年ごとの推移
    rows = [a for a in annual if a["病気"] == name]
    metrics = list(dict.fromkeys(a["指標"] for a in rows))
    years = sorted({a["年"] for a in rows})
    r = section(ws, r, "年ごとの推移")
    r = header(ws, r, ["年"] + metrics + [f"前年比（{main}）", "区分", "出典ID"])
    first = r
    main_col = metrics.index(main) + 2
    by = {(a["年"], a["指標"]): a for a in rows}
    for y in years:
        status = by[(y, metrics[0])]["区分"]
        fill = fill_soku if status == "速報" else None
        put(ws, r, 1, y, f_bold, fill, '0"年"', Alignment(horizontal="left"))
        for j, m in enumerate(metrics, 2):
            a = by.get((y, m))
            if name == "HIV/エイズ" and m == "合計":
                # 合計は感染者と患者の和（2016〜2022年は資料に合計がないため）
                put(ws, r, j, f"=B{r}+C{r}", fill=fill, fmt="#,##0")
            else:
                put(ws, r, j, a["値"] if a else None, fill=fill, fmt=numfmt(m, a["値"] if a else 0))
        cl = get_column_letter(main_col)
        put(ws, r, len(metrics) + 2, f'=IF({cl}{r - 1}="","",{cl}{r}/{cl}{r - 1}-1)' if r > first else "—",
            fill=fill, fmt='+0.0%;-0.0%;0.0%', align=Alignment(horizontal="right"))
        put(ws, r, len(metrics) + 3, status, fill=fill, align=Alignment(horizontal="center"))
        put(ws, r, len(metrics) + 4, by[(y, metrics[0])]["出典ID"], fill=fill)
        r += 1
    confirmed = [y for y in years if by[(y, metrics[0])]["区分"] == "確定"]
    row_of = {y: first + i for i, y in enumerate(years)}
    cl = get_column_letter(main_col)
    links[name] = {"latest_year": f"'{sheet}'!A{row_of[confirmed[-1]]}", "latest": f"'{sheet}'!{cl}{row_of[confirmed[-1]]}",
                   "prev": f"'{sheet}'!{cl}{row_of[confirmed[-2]]}", "first_year": f"'{sheet}'!A{row_of[confirmed[0]]}",
                   "first": f"'{sheet}'!{cl}{row_of[confirmed[0]]}",
                   "soku": f"'{sheet}'!{cl}{row_of[years[-1]]}" if years[-1] not in confirmed else None}
    r = para(ws, r, "黄色の行は速報値（あとで変わることがあります）。それ以外は確定値。前年比はこの表で計算したもの。", f_note)
    r += 1

    # 2026年の途中経過
    r = section(ws, r, "2026年の途中経過（速報）")
    r = header(ws, r, ["指標", "期間", "2026年", "前年同期", "前年同期比", "区分", "出典ID"])
    for t in [t for t in ytd if t["病気"] == name]:
        put(ws, r, 1, t["指標"], align=wrap)
        put(ws, r, 2, t["期間"], align=wrap)
        put(ws, r, 3, t["2026年"], fmt="#,##0")
        put(ws, r, 4, t["前年同期"], fmt="#,##0")
        put(ws, r, 5, f"=C{r}/D{r}-1", fmt='+0.0%;-0.0%;0.0%')
        put(ws, r, 6, t["区分"], align=Alignment(horizontal="center"))
        put(ws, r, 7, t["出典ID"])
        ws.row_dimensions[r].height = 50
        if t["指標"] == ymain:
            links[name].update(ytd_label=f"'{sheet}'!B{r}", ytd=f"'{sheet}'!C{r}", ytd_prev=f"'{sheet}'!D{r}")
        r += 1
    detail = next(t for t in ytd if t["病気"] == name)["詳しい期間"]
    r = para(ws, r, "期間のくわしい説明：" + detail, f_note)
    for n in notes_of(name, "表の注", "2026年の途中経過"):
        r = para(ws, r, "注：" + n["文"], f_note)
    r += 1

    # 内訳の表
    rows = [b for b in breakdown if b["病気"] == name and b["種類"] != "公費"]
    for title in dict.fromkeys(b["表"] for b in rows):
        t_rows = [b for b in rows if b["表"] == title]
        cols = list(dict.fromkeys(b["列"] for b in t_rows))
        r = section(ws, r, title)
        r = header(ws, r, ["区分"] + cols + ["出典ID"])
        cell = {(b["区分"], b["列"]): b for b in t_rows}
        for label in dict.fromkeys(b["区分"] for b in t_rows):
            put(ws, r, 1, label, align=wrap)
            for j, c in enumerate(cols, 2):
                v = cell[(label, c)]["値"] if (label, c) in cell else None
                if v is None:
                    put(ws, r, j, "資料になし", f_note, align=Alignment(horizontal="right"))
                else:
                    put(ws, r, j, v, fmt="0" if c == "順位" else numfmt(c, v))
            put(ws, r, len(cols) + 2, t_rows[0]["出典ID"])
            if len(str(label)) > 22:
                ws.row_dimensions[r].height = 30
            r += 1
        for n in notes_of(name, "表の注", title):
            r = para(ws, r, "注：" + n["文"], f_note)
        r += 1

    for kind, head in [("特徴", "最近の特徴"), ("対策", "国の主な対策"), ("公費", "公費で受けられる支援")]:
        r = section(ws, r, head)
        for n in notes_of(name, kind):
            r = para(ws, r, f"{n['番号']}. {n['文']}（出典 {n['出典ID']}）")
        if kind == "公費":
            r = para(ws, r, "4つの病気の一覧や、HIVの治療の自己負担の上限の表は「公費の支援」シートにあります。", f_note)
        r += 1
    r = section(ws, r, "この病気の数字を見るときの注意")
    for n in notes_of(name, "見るときの注意"):
        r = para(ws, r, f"・{n['文']}")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

# HIV/エイズの世界の動向（UNAIDSの推計）
wsw = wb.create_sheet("HIV・世界の動向", wb.sheetnames.index("HIV・エイズ") + 1)
wsw.sheet_view.showGridLines = False
wsw.cell(row=1, column=1, value="HIV/エイズ 世界の動向").font = f_title
wsw.row_dimensions[1].height = 26
r = para(wsw, 2, "数え方：" + notes_of(WORLD, "数え方")[0]["文"], f_note, cols=13, per_line=110)
r += 1
wcols = ["区分", "協力覚書", "感染している人", "感染している人（下限）", "感染している人（上限）", "15〜49歳の感染率（%）", "新たな感染",
         "新たな感染の増減（2010年比・%）", "エイズによる死亡", "死亡の増減（2010年比・%）", "感染を知っている人（%）",
         "治療を受けている人（%）", "ウイルスが抑えられている人（%）"]
for part, rows_w in [(f"世界と地域（{world[0]['年']}年）", [w for w in world if w["区分"] != "国"]),
                     (f"日本と協力覚書を結んだ国（{world[0]['年']}年、感染している人の多い順）",
                      sorted([w for w in world if w["区分"] == "国"], key=lambda w: -(w["感染している人"] if isinstance(w["感染している人"], (int, float)) else -1)))]:
    r = section(wsw, r, part)
    r = header(wsw, r, ["名前"] + wcols + ["出典ID"])
    for w in rows_w:
        put(wsw, r, 1, w["名前"], f_bold, fill_soft)
        for j, c in enumerate(wcols, 2):
            v = w[c]
            if v is None:
                put(wsw, r, j, "—" if c != "協力覚書" else "", f_note, align=Alignment(horizontal="right"))
            elif isinstance(v, str):
                put(wsw, r, j, v, align=Alignment(horizontal="right" if c not in ("区分", "協力覚書") else "left", wrap_text=True))
            else:
                put(wsw, r, j, v, fmt='+0;-0;0' if "増減" in c else ("0.0" if isinstance(v, float) else "#,##0"))
        put(wsw, r, len(wcols) + 2, w["出典ID"])
        r += 1
    r += 1
r = para(wsw, r, "「<100」は100人未満、「—」はUNAIDSに数字がないことを表します。増減（2010年比・%）と割合（%）はUNAIDSの資料に書かれた値です。", f_note, cols=13, per_line=110)
r += 1
for kind, head in [("協力覚書", "日本の協力覚書（特定技能・技能実習・育成就労）"), ("移住者とHIV", "移住者とHIV（国連の機関の見方）"),
                   ("入国前の健康診断", "日本の入国前の健康診断とIOM")]:
    r = section(wsw, r, head)
    for n in notes_of(WORLD, kind):
        r = para(wsw, r, f"{n['番号']}. {n['文']}" + (f"（出典 {n['出典ID']}）" if n["出典ID"] else ""), cols=13, per_line=110)
    r += 1
r = section(wsw, r, "この数字を見るときの注意")
for n in notes_of(WORLD, "見るときの注意"):
    r = para(wsw, r, f"・{n['文']}", cols=13, per_line=110)
wsw.column_dimensions["A"].width = 26
for i, w in enumerate([8, 22, 13, 12, 12, 12, 12, 13, 12, 13, 12, 12, 13, 10], 2):
    wsw.column_dimensions[get_column_letter(i)].width = w
wsw.freeze_panes = "B4"
wsw.page_setup.orientation = "landscape"
wsw.page_setup.fitToWidth = 1
wsw.page_setup.fitToHeight = 0
wsw.sheet_properties.pageSetUpPr.fitToPage = True

# 4つの比較
sm = wb.create_sheet("4つの比較", 1)
sm.sheet_view.showGridLines = False
sm.cell(row=1, column=1, value="4つの病気の比較").font = f_title
para(sm, 2, "数字は各病気のシートから自動で読み込んでいます。前年比・10年前と比べた倍率・前年同期比はここで計算したものです。", f_note, cols=14)
heads = ["病気", "数えているもの", "最新の確定値の年", "最新の確定値", "その前の年", "前年比",
         "10年前の年", "10年前の数", "10年前と比べて", "2025年（速報値）", "2026年の途中経過の期間",
         "2026年の途中経過", "前年の同じ時期", "前年同期比"]
header(sm, 4, heads)
widths = [18, 30, 12, 13, 13, 11, 11, 12, 12, 13, 30, 13, 13, 11]
for i, w in enumerate(widths, 1):
    sm.column_dimensions[get_column_letter(i)].width = w
r = 5
for name, sheet, display, main, ymain in SHEETS:
    L = links[name]
    put(sm, r, 1, display, f_bold, fill_soft)
    put(sm, r, 2, WHAT[name] if name != "結核" else WHAT[name], align=wrap)
    put(sm, r, 3, f"={L['latest_year']}", fmt='0"年"')
    put(sm, r, 4, f"={L['latest']}", f_bold, fmt="#,##0")
    put(sm, r, 5, f"={L['prev']}", fmt="#,##0")
    put(sm, r, 6, f"=D{r}/E{r}-1", fmt='+0.0%;-0.0%;0.0%')
    put(sm, r, 7, f"={L['first_year']}", fmt='0"年"')
    put(sm, r, 8, f"={L['first']}", fmt="#,##0")
    put(sm, r, 9, f"=D{r}/H{r}", fmt='0.0"倍"')
    put(sm, r, 10, f"={L['soku']}" if L["soku"] else "—", fmt="#,##0", fill=fill_soku if L["soku"] else None,
        align=None if L["soku"] else Alignment(horizontal="center"))
    put(sm, r, 11, f"={L['ytd_label']}", align=wrap)
    put(sm, r, 12, f"={L['ytd']}", fmt="#,##0", fill=fill_soku)
    put(sm, r, 13, f"={L['ytd_prev']}", fmt="#,##0", fill=fill_soku)
    put(sm, r, 14, f"=L{r}/M{r}-1", fmt='+0.0%;-0.0%;0.0%', fill=fill_soku)
    sm.row_dimensions[r].height = 36
    r += 1
r += 1
for t in ["・最新の確定値が出ている年は病気ごとに違います（HIV/エイズと結核は2025年、梅毒とはしかは2024年）。",
          "・「10年前の年」は、最新の確定値の年から数えて10年分の最初の年です。",
          "・2026年の途中経過は速報値で、病気によって期間が違います。結核の途中経過は医師の届出の数で、発病していない感染者も含むため、年報の「新登録患者数」とは比べられません。",
          "・黄色のセルは速報値です。"]:
    r = para(sm, r, t, f_note, cols=14, per_line=120)
sm.freeze_panes = "B5"
sm.page_setup.orientation = "landscape"
sm.sheet_properties.pageSetUpPr.fitToPage = True
sm.page_setup.fitToHeight = 0

# 公費で受けられる支援
ks = wb.create_sheet("公費の支援", wb.sheetnames.index("4つの比較") + 1)
ks.sheet_view.showGridLines = False
ks.cell(row=1, column=1, value="公費で受けられる支援（検査・予防・治療）").font = f_title
ks.row_dimensions[1].height = 26
for i, w in enumerate([16, 13, 16, 71, 11], 1):
    ks.column_dimensions[get_column_letter(i)].width = w
r = para(ks, 2, notes_of(KOUHI, "説明")[0]["文"], f_note, cols=5, per_line=95)
r += 1
r = section(ks, r, "ひと目でわかる一覧")
r = header(ks, r, ["病気", "場面", "目安", "内容", "出典ID"])
first_k = r
levels = [n for n in notes if n["病気"] == KOUHI and n["種類"] == "目安"]
level_fill = dict(zip([n["表"] for n in levels], [fill_yes, fill_cond, fill_no]))
for k in kouhi:
    yes = level_fill[k["目安"]] is not fill_no
    put(ks, r, 1, "はしか（麻しん）" if k["病気"] == "はしか" else k["病気"], f_bold, fill_soft, align=Alignment(vertical="top"))
    put(ks, r, 2, k["場面"], align=Alignment(vertical="top"))
    put(ks, r, 3, k["目安"], f_bold if yes else f_body, level_fill[k["目安"]], align=Alignment(horizontal="center", vertical="top", wrap_text=True))
    put(ks, r, 4, k["内容"], align=wrap)
    put(ks, r, 5, k["出典ID"], align=Alignment(vertical="top"))
    ks.row_dimensions[r].height = max(18, 16 * math.ceil(len(k["内容"]) / 41), 16 * math.ceil(len(k["目安"]) / 7))
    r += 1
ks.auto_filter.ref = f"A{first_k - 1}:E{r - 1}"
for n in levels:
    r = para(ks, r, f"{n['印']}「{n['表']}」…{n['文']}", f_note, cols=5, per_line=95)
r += 1
r = section(ks, r, "治療費の公費のしくみ（感染症法）")
for n in notes_of(KOUHI, "しくみ"):
    r = para(ks, r, f"{n['番号']}. {n['文']}（出典 {n['出典ID']}）", cols=5, per_line=95)
r += 1
for name, sheet, display, main, ymain in SHEETS:
    r = section(ks, r, display)
    for n in notes_of(name, "公費"):
        r = para(ks, r, f"{n['番号']}. {n['文']}（出典 {n['出典ID']}）", cols=5, per_line=95)
    for title in dict.fromkeys(b["表"] for b in breakdown if b["病気"] == name and b["種類"] == "公費"):
        t_rows = [b for b in breakdown if b["病気"] == name and b["表"] == title]
        r += 1
        ks.cell(row=r, column=1, value=title).font = f_bold
        r += 1
        ks.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        put(ks, r, 1, "所得の区分（健康保険の世帯）", f_head, fill_head, align=center)
        put(ks, r, 4, t_rows[0]["列"], f_head, fill_head, align=center)
        put(ks, r, 5, "出典ID", f_head, fill_head, align=center)
        r += 1
        for b in t_rows:
            ks.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            put(ks, r, 1, b["区分"], align=wrap)
            put(ks, r, 4, b["値"], fmt='#,##0"円"', align=Alignment(horizontal="left", vertical="top"))
            put(ks, r, 5, b["出典ID"])
            ks.row_dimensions[r].height = 30 if len(b["区分"]) > 26 else 18
            r += 1
        for n in notes_of(name, "表の注", title):
            r = para(ks, r, "注：" + n["文"], f_note, cols=5, per_line=95)
    r += 1
r = section(ks, r, "この内容を見るときの注意")
for n in notes_of(KOUHI, "見るときの注意"):
    r = para(ks, r, f"・{n['文']}", cols=5, per_line=95)
ks.freeze_panes = "A5"
ks.page_setup.orientation = "landscape"
ks.page_setup.fitToWidth = 1
ks.page_setup.fitToHeight = 0
ks.sheet_properties.pageSetUpPr.fitToPage = True


# ===== 国別の比較（WHO）=====
def sheet(title, idx=None):
    w = wb.create_sheet(title, idx) if idx is not None else wb.create_sheet(title)
    w.sheet_view.showGridLines = False
    return w


def finish(w, freeze=None, landscape=True):
    if freeze:
        w.freeze_panes = freeze
    if landscape:
        w.page_setup.orientation = "landscape"
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True


pop_of = {}
for r0 in json.load(open(os.path.join(os.path.dirname(os.path.abspath(data_dir)), "research", "world_stats.json"), encoding="utf-8"))["rows"]:
    if r0.get("tb"):
        pop_of[r0["iso3"]] = r0["tb"].get("pop")
cw = sheet("国別の比較", 1)
cw.cell(row=1, column=1, value="国別の比較（WHOの公開データ）・協力覚書・入国前の結核検査").font = f_title
cw.row_dimensions[1].height = 26
r = para(cw, 2, "数え方：" + notes_of("国別の比較", "数え方")[0]["文"], f_note, cols=16, per_line=140)
r = para(cw, r, "その国に住む人全体（一般人口）の数字で、その国から日本へ入国する人の感染の割合ではありません。はしかの「人口100万人あたり」は、報告数をWHOの結核推計の人口で割ってこの表で計算しています（式が入っています）。最初は日本と関係の深い国（協力覚書・入国前の結核検査の対象国と日本）を並べ、そのあとにほかの国・地域を並べています。", f_note, cols=16, per_line=140)
r += 1
ccols = [("国", "国"), ("英語名", "英語名"), ("地域", "地域"), ("協力覚書", "協力覚書"), ("入国前結核検査", "入国前の結核検査（JPETS）"), ("入国前結核検査_開始", "JPETSの開始"),
         ("結核_年", "結核：年"), ("結核_罹患率", "結核：罹患率（人口10万対・推計）"), ("結核_罹患率_下限", "罹患率の下限"), ("結核_罹患率_上限", "罹患率の上限"),
         ("結核_人数", "結核：人数（推計）"), ("結核_前年の罹患率", "結核：前年の罹患率"),
         ("はしか_年", "はしか：年"), ("はしか_報告数", "はしか：報告数"), ("はしか_前年の報告数", "はしか：前年の報告数"), ("人口", "人口（WHO結核推計の人口）"), ("はしか_100万人あたり", "はしか：人口100万人あたり（計算）"),
         ("HIV_年", "HIV：年"), ("HIV_15〜49歳の感染率", "HIV：15〜49歳の感染率（%）"), ("HIV_感染している人", "HIV：感染している人（推計）"), ("HIV_新たな感染", "HIV：新たな感染（推計）"),
         ("梅毒_妊婦の陽性率", "梅毒：妊婦健診の陽性率（%）"), ("梅毒_妊婦の陽性率_年", "その年"), ("梅毒_報告数", "梅毒：報告数"), ("梅毒_報告数_年", "その年"),
         ("梅毒_先天梅毒の率", "梅毒：先天梅毒（出生10万対）"), ("梅毒_先天梅毒の率_年", "その年")]
TEXT_COLS = ("国", "英語名", "地域", "協力覚書", "入国前結核検査", "入国前結核検査_開始")
r = header(cw, r, [t for _, t in ccols])
cw.row_dimensions[r - 1].height = 48
first_c = r
order = sorted(countries, key=lambda c: (not c["関係の深い国"], c["国"]))
col_of = {k: i for i, (k, _) in enumerate(ccols, 1)}
for c in order:
    for j, (k, _) in enumerate(ccols, 1):
        if k == "人口":
            v = pop_of.get(c["ISO3"])
        elif k == "はしか_100万人あたり":
            v = (f"={get_column_letter(col_of['はしか_報告数'])}{r}/{get_column_letter(col_of['人口'])}{r}*1000000"
                 if c.get("はしか_報告数") is not None and pop_of.get(c["ISO3"]) else None)
        else:
            v = c.get(k)
        fill = fill_soft if k == "国" and c["関係の深い国"] else None
        if v is None or v == "":
            put(cw, r, j, "" if k in TEXT_COLS else "—", f_note, fill=fill, align=Alignment(horizontal="center"))
        elif isinstance(v, str) and not v.startswith("="):
            put(cw, r, j, v, f_bold if k == "国" else f_body, fill=fill, align=Alignment(horizontal="left" if k in TEXT_COLS else "right", wrap_text=True))
        else:
            if isinstance(v, float) and v == int(v) and k != "はしか_100万人あたり":
                v = int(v)
            fm = '0"年"' if k.endswith("年") else "0.0" if (k == "はしか_100万人あたり" or isinstance(v, float)) else "#,##0"
            put(cw, r, j, v, fmt=fm)
    r += 1
cw.auto_filter.ref = f"A{first_c - 1}:{get_column_letter(len(ccols))}{r - 1}"
r += 1
r = section(cw, r, "この数字を見るときの注意")
for n in notes_of("国別の比較", "見るときの注意"):
    r = para(cw, r, f"・{n['文']}", cols=16, per_line=140)
r = para(cw, r, "出典：WHO-1（結核）、WHO-2（はしか）、WHO-3（HIV）、WHO-4（梅毒）、WHO-5（地域の区分）。資料名とURLは「出典一覧」シートにあります。", f_note, cols=16, per_line=140)
for i, (k, _) in enumerate(ccols, 1):
    cw.column_dimensions[get_column_letter(i)].width = {"国": 16, "英語名": 22, "地域": 12, "協力覚書": 20, "入国前結核検査": 11, "入国前結核検査_開始": 11}.get(k, 11)
finish(cw, f"B{first_c}")

# ===== 協力覚書・入国前対策 =====
ST_FILL = {"実施中": fill_yes, "開始予定": fill_cond, "未確認": fill_cond}
ew = sheet("協力覚書・入国前対策", 2)
ew.cell(row=1, column=1, value="日本の入国前対策（協力覚書と、入国前・入国時の健康確認）").font = f_title
ew.row_dimensions[1].height = 26
for i, w in enumerate([16, 12, 46, 14, 16, 60, 40, 16], 1):
    ew.column_dimensions[get_column_letter(i)].width = w
r = para(ew, 2, "協力覚書（MOU・MOC）を結んでいることと、入国前の健康診断・感染症検査が義務になっていることは、別の制度です。状態は「実施中」「開始予定」「制度対象外」「検査制度なし」「未確認」で区別しています（資料が見つからないものは「未確認」とし、制度がないと決めつけていません）。", f_note, cols=8, per_line=150)
r += 1
r = section(ew, r, "4つの感染症 × 入国前・入国時の制度")
r = header(ew, r, ["病気", "状態", "制度", "", "", "説明", "", "出典ID"])
for x in [e for e in entry if e["区分"] == "病気と制度"]:
    put(ew, r, 1, x["病気"], f_bold, fill_soft, align=wrap)
    put(ew, r, 2, x["状態"], f_bold, ST_FILL.get(x["状態"], fill_no), align=center)
    ew.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    put(ew, r, 3, x["名前"], align=wrap)
    ew.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    put(ew, r, 6, x.get("内容"), align=wrap)
    put(ew, r, 8, x.get("出典ID"), align=wrap)
    ew.row_dimensions[r].height = max(20, 15 * math.ceil(len(x.get("内容") or "") / 60))
    r += 1
r += 1
r = section(ew, r, "入国前結核スクリーニング（JPETS）の対象国")
r = header(ew, r, ["国", "状態", "内容", "開始", "", "", "", "出典ID"])
for x in [e for e in entry if e["区分"] == "JPETSの国"]:
    put(ew, r, 1, x["国"], f_bold, fill_soft)
    put(ew, r, 2, x["状態"], f_bold, ST_FILL.get(x["状態"], fill_no), align=center)
    put(ew, r, 3, x.get("内容"), align=wrap)
    put(ew, r, 4, x.get("開始"), align=wrap)
    put(ew, r, 8, x.get("出典ID"), align=wrap)
    ew.row_dimensions[r].height = max(20, 15 * math.ceil(len(x.get("内容") or "") / 38))
    r += 1
r += 1
r = section(ew, r, "入国前・入国時の制度")
LBL = {"内容": "制度", "検査の内容": "検査・健康診断の内容", "対象": "対象", "対象外": "対象外", "開始": "開始", "有効期間": "有効期間", "根拠": "根拠"}
for x in [e for e in entry if e["区分"] == "制度"]:
    ew.cell(row=r, column=1, value=f"{x['名前']}（{x['状態']}）").font = f_bold
    r += 1
    for k, lbl in LBL.items():
        if x.get(k):
            r = para(ew, r, f"{lbl}：{x[k]}", cols=8, per_line=150)
    r = para(ew, r, f"出典：{x.get('出典ID') or '—'}", f_note, cols=8, per_line=150)
    r += 1
r = section(ew, r, "在留資格ごとの適用範囲")
r = header(ew, r, ["在留資格", "状態", "JPETS", "", "", "入国前の健康診断など", "", "出典ID"])
for x in [e for e in entry if e["区分"] == "在留資格"]:
    put(ew, r, 1, x["名前"], f_bold, fill_soft)
    put(ew, r, 2, x["状態"], align=center)
    ew.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    put(ew, r, 3, x.get("JPETS"), align=wrap)
    ew.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    put(ew, r, 6, x.get("内容"), align=wrap)
    put(ew, r, 8, x.get("出典ID"), align=wrap)
    ew.row_dimensions[r].height = max(20, 15 * math.ceil(max(len(x.get("内容") or ""), len(x.get("JPETS") or "") * 1.5) / 60))
    r += 1
r += 1
r = section(ew, r, "協力覚書（特定技能・技能実習・育成就労）")
r = header(ew, r, ["国", "制度", "正式名称", "日付", "日付の種類", "協力の目的", "資料URL", "出典ID"])
for x in [e for e in entry if e["区分"] == "協力覚書"]:
    put(ew, r, 1, x["国"], f_bold, fill_soft)
    put(ew, r, 2, x["制度"])
    put(ew, r, 3, x.get("名前"), align=wrap)
    put(ew, r, 4, x.get("日付"))
    put(ew, r, 5, x.get("日付の種類"), align=wrap)
    put(ew, r, 6, x.get("内容"), align=wrap)
    c = put(ew, r, 7, x.get("資料URL"), f_link, align=wrap)
    if x.get("資料URL"):
        c.hyperlink = x["資料URL"]
    put(ew, r, 8, x.get("出典ID"), align=wrap)
    ew.row_dimensions[r].height = max(30, 15 * math.ceil(max(len(x.get("名前") or "") / 22, len(x.get("内容") or "") / 30)))
    r += 1
r += 1
r = section(ew, r, "協力覚書の本文にある健康・感染症の記載")
for x in [e for e in entry if e["区分"] == "覚書の健康の記載"]:
    r = para(ew, r, f"{x['国']}：{x.get('内容') or '—'}", cols=8, per_line=150)
r += 1
r = section(ew, r, "協力覚書についての説明")
for n in notes_of(WORLD, "協力覚書") + notes_of("入国前対策", "協力覚書"):
    r = para(ew, r, f"・{n['文']}" + (f"（出典 {n['出典ID']}）" if n["出典ID"] else ""), cols=8, per_line=150)
r += 1
r = section(ew, r, "制度上の確認事項")
for x in [e for e in entry if e["区分"] == "確認事項"]:
    r = para(ew, r, f"【{x['名前']}】{x['内容']}（出典 {x.get('出典ID') or '—'}）", cols=8, per_line=150)
finish(ew, "A4")

# ===== 医療費 =====
if costs:
    mw = sheet("医療費", wb.sheetnames.index("公費の支援") + 1 if "公費の支援" in wb.sheetnames else None)
    mw.cell(row=1, column=1, value="感染症の治療には、1人当たりいくらかかる？（医療費の総額と、税金・公的医療保険・本人負担）").font = f_title
    mw.row_dimensions[1].height = 26
    for i, w in enumerate([12, 30, 16, 12, 16, 14, 34, 34, 24, 40, 14], 1):
        mw.column_dimensions[get_column_letter(i)].width = w
    r = para(mw, 2, "「医療費の総額」は「税金の負担」と同じではありません。公表値を確認できなかった項目は「公表値を確認できず」とし、推測の数字は入れていません。「制度上の額」は国が決めた価格（薬価・点数・上限額）、「試算」は制度のきまりから計算したもので、実際に払われた額（実績値）とは違います。", f_note, cols=11, per_line=170)
    r += 1
    r = header(mw, r, ["病気", "項目", "金額など", "単位", "対象年・時点", "数字の種類", "対象", "範囲", "負担", "注・計算式・探した資料", "出典ID"])
    first_m = r
    for x in [c for c in costs if c["区分"] == "費用"]:
        put(mw, r, 1, x["病気"], f_bold, fill_soft, align=wrap)
        put(mw, r, 2, x.get("名前") or x["項目"], align=wrap)
        v = x.get("値")
        if v is None:
            put(mw, r, 3, x.get("状態") or "公表値を確認できず", f_note, fill_no, align=wrap)
        else:
            put(mw, r, 3, v, f_bold, fmt="#,##0.##" if isinstance(v, (int, float)) else None, align=wrap if isinstance(v, str) else None)
        put(mw, r, 4, x.get("単位"), align=wrap)
        put(mw, r, 5, x.get("年"), align=wrap)
        put(mw, r, 6, x.get("種類"), align=wrap)
        put(mw, r, 7, x.get("対象"), align=wrap)
        put(mw, r, 8, x.get("範囲"), align=wrap)
        put(mw, r, 9, x.get("負担"), align=wrap)
        put(mw, r, 10, "／".join(t for t in [x.get("注"), f"計算：{x['計算式']}" if x.get("計算式") else None, f"探した資料：{x['探した資料']}" if x.get("探した資料") and v is None else None] if t), align=wrap)
        put(mw, r, 11, x.get("出典ID"), align=wrap)
        mw.row_dimensions[r].height = max(30, 15 * math.ceil(max(len(x.get("対象") or "") / 18, len(x.get("範囲") or "") / 18, len(x.get("注") or "") / 22 + (len(x.get("探した資料") or "") / 22 if v is None else 0))))
        r += 1
    mw.auto_filter.ref = f"A{first_m - 1}:K{r - 1}"
    r += 1
    HEAD = {"財源": "参考：医療費の財源", "保険のきまり": "参考：公的医療保険のきまり", "試算のきまり": "試算のきまり"}
    for x in [c for c in costs if c["区分"] in HEAD]:
        r = section(mw, r, f"{HEAD[x['区分']]}（{x.get('項目')}）")
        r = para(mw, r, f"{x.get('名前') or ''} {x.get('値') or ''}（{x.get('年') or ''}）{x.get('注') or ''}（出典 {x.get('出典ID') or '—'}）", cols=11, per_line=170)
    r += 1
    r = section(mw, r, "この内容を見るときの注意")
    for n in notes_of("医療費", "比べ方") + notes_of("医療費", "試算について"):
        r = para(mw, r, f"・{n['文']}", cols=11, per_line=170)
    finish(mw, f"C{first_m}")

# ===== 病気の解説 =====
if info:
    iw = sheet("病気の解説")
    iw.cell(row=1, column=1, value="4つの病気の解説（概要・感染経路・世界の状況・検査・予防・治療）").font = f_title
    for i, w in enumerate([12, 16, 30, 90, 12, 10, 12, 12, 16], 1):
        iw.column_dimensions[get_column_letter(i)].width = w
    SEC = {"overview": "概要", "routes": "感染経路", "not_routes": "日常生活ではうつらないこと", "world": "世界の状況", "world_numbers": "世界の数字", "test": "検査", "prevent": "予防", "treat": "治療", "special": "特記事項", "special_numbers": "特記事項の数字"}
    r = header(iw, 3, ["病気", "項目", "見出し", "内容", "値", "単位", "年", "種類", "出典ID"])
    for x in info:
        put(iw, r, 1, x["病気"], f_bold, fill_soft)
        put(iw, r, 2, SEC.get(x["節"], x["節"]))
        put(iw, r, 3, x.get("見出し"), align=wrap)
        put(iw, r, 4, x.get("文"), align=wrap)
        put(iw, r, 5, x.get("値"), fmt="#,##0.##" if isinstance(x.get("値"), (int, float)) else None)
        put(iw, r, 6, x.get("単位"))
        put(iw, r, 7, x.get("年"))
        put(iw, r, 8, x.get("種類"), align=wrap)
        put(iw, r, 9, x.get("出典ID"), align=wrap)
        iw.row_dimensions[r].height = max(18, 15 * math.ceil(len(x.get("文") or "") / 58))
        r += 1
    iw.auto_filter.ref = f"A3:I{r - 1}"
    finish(iw, "A4")

# ===== その他の感染症・さらに詳しく調べたい方へ =====
if others or guide:
    ow = sheet("その他・資料の案内")
    ow.cell(row=1, column=1, value="その他の感染症（補足）と、さらに詳しく調べたい方への資料").font = f_title
    for i, w in enumerate([18, 50, 40, 50, 16, 40], 1):
        ow.column_dimensions[get_column_letter(i)].width = w
    r = 3
    if others:
        r = section(ow, r, "その他の感染症（補足）")
        r = header(ow, r, ["病名", "病気の特徴", "主な感染経路", "予防・対策", "出典ID", ""])
        for o in others:
            put(ow, r, 1, o["病名"], f_bold, fill_soft)
            put(ow, r, 2, o.get("特徴"), align=wrap)
            put(ow, r, 3, o.get("感染経路"), align=wrap)
            put(ow, r, 4, o.get("予防・対策"), align=wrap)
            put(ow, r, 5, o.get("出典ID"), align=wrap)
            ow.row_dimensions[r].height = max(30, 15 * math.ceil(max(len(o.get("特徴") or ""), len(o.get("予防・対策") or "")) / 30))
            r += 1
        r += 1
    if guide:
        r = section(ow, r, "さらに詳しく調べたい方へ")
        r = header(ow, r, ["分類", "資料名", "発行機関", "調べられる内容", "公表日・対象期間", "URL"])
        for g in guide:
            put(ow, r, 1, g["分類"], f_bold, fill_soft, align=wrap)
            put(ow, r, 2, g["資料名"], align=wrap)
            put(ow, r, 3, g["発行機関"], align=wrap)
            put(ow, r, 4, g.get("調べられる内容"), align=wrap)
            put(ow, r, 5, f"{g.get('公表日') or '記載なし'}／{g.get('対象期間') or '—'}", align=wrap)
            c = put(ow, r, 6, g["URL"], f_link, align=wrap)
            c.hyperlink = g["URL"]
            ow.row_dimensions[r].height = max(30, 15 * math.ceil(max(len(g.get("調べられる内容") or "") / 30, len(g["URL"]) / 38)))
            r += 1
    finish(ow, "A3")

# 見るときの注意
cs = wb.create_sheet("見るときの注意")
cs.sheet_view.showGridLines = False
cs.column_dimensions["A"].width = 16
cs.column_dimensions["B"].width = 110
cs.cell(row=1, column=1, value="数字を見るときの注意").font = f_title
r = 3
for who in ["全体", "国別の比較", "HIV/エイズ", WORLD, "梅毒", "はしか", "結核", KOUHI]:
    cs.cell(row=r, column=1, value={"全体": "全体にかかわること", KOUHI: "公費で受けられる支援", "国別の比較": "国別の比較（WHO）"}.get(who, who)).font = f_sec
    r += 1
    for n in [n for n in notes if n["病気"] == who and n["種類"] == "見るときの注意"]:
        put(cs, r, 1, n["番号"], align=Alignment(horizontal="center", vertical="top"), border=False)
        c = put(cs, r, 2, n["文"], align=wrap, border=False)
        cs.row_dimensions[r].height = max(18, 16 * math.ceil(len(n["文"]) / 70))
        r += 1
    r += 1

# 出典一覧
ss = wb.create_sheet("出典一覧")
ss.sheet_view.showGridLines = False
ss.cell(row=1, column=1, value="出典一覧（国・自治体・国の研究機関・国連機関の公式資料のみ）").font = f_title
para(ss, 2, "各シートの「出典ID」はこの表の番号です。URLはクリックで開けます。確認日は2026年10月26日。", f_note, cols=7)
header(ss, 4, ["出典ID", "病気", "資料名", "発行元", "URL", "公表日・更新日", "確認日"])
for i, w in enumerate([10, 14, 60, 30, 55, 22, 12], 1):
    ss.column_dimensions[get_column_letter(i)].width = w
r = 5
for s in sources:
    put(ss, r, 1, s["出典ID"])
    put(ss, r, 2, s["病気"])
    put(ss, r, 3, s["資料名"], align=wrap)
    put(ss, r, 4, s["発行元"], align=wrap)
    c = put(ss, r, 5, s["URL"], f_link, align=wrap)
    c.hyperlink = s["URL"]
    put(ss, r, 6, s["公表日・更新日"], align=wrap)
    put(ss, r, 7, s["確認日"])
    ss.row_dimensions[r].height = max(18, 15 * math.ceil(max(len(s["資料名"]) / 30, len(s["URL"]) / 50)))
    r += 1
ss.freeze_panes = "A5"
ss.auto_filter.ref = f"A4:G{r - 1}"

# はじめに
intro.sheet_view.showGridLines = False
intro.column_dimensions["A"].width = 22
intro.column_dimensions["B"].width = 100
intro.cell(row=1, column=1, value="世界の感染症と日本の水際対策：HIV/エイズ・梅毒・はしか・結核").font = f_title
rows = [
    ("作成日", "2026年10月26日（この日までに国が公表した資料をもとにしています）"),
    ("内容", "4つの感染症について、国別の感染状況（WHO）、日本と協力覚書を結んだ国と入国前の健康確認の制度、国内の最新の人数・過去10年の動き・男女・年代・都道府県・2026年の途中経過、国の対策、検査・予防・治療で使える公費の制度、治療費と公的負担をまとめました。国別の数字はその国の人口全体の数字で、日本へ入国する人の数字ではありません。"),
    ("使った資料", "国内の数字と制度は国の公式資料（厚生労働省、国立健康危機管理研究機構〔JIHS〕、出入国在留管理庁、外国人技能実習機構、e-Gov法令検索など）を使いました。HIV/エイズの年報は、厚生労働省のページから案内されている「エイズ予防情報ネット」（国の委託先）に置かれているものを使っています。国別の数字と世界の動向は、国連の機関（世界保健機関〔WHO〕・国連合同エイズ計画〔UNAIDS〕・国際移住機関〔IOM〕・国連統計部）の公式資料を使っています。公費の制度のうち結核の医療費助成の一部などは、自治体の公式ページで確かめました。"),
    ("シートの説明", "「国別の比較」…世界の国・地域の結核・はしか・HIV・梅毒の数字（WHO）と、協力覚書・入国前の結核検査。\n「協力覚書・入国前対策」…協力覚書の一覧と、入国前・入国時の健康確認の制度、制度上の確認事項。\n「4つの比較」…4つの病気を1つの表で比べます。\n「公費の支援」…検査・予防・治療で使える公費の制度の一覧と、病気ごとの説明、HIVの治療の自己負担の上限。\n「医療費」…治療費と、税金・公的医療保険・本人負担の内訳（確認できた公表値と、確認できなかった項目）。\n「HIV・エイズ」「梅毒」「はしか」「結核」…病気ごとの詳しい数字です。\n「HIV・世界の動向」…世界全体と、日本と協力覚書を結んだ国のHIVの状況（UNAIDSの推計）、移住者とHIV、入国前の健康診断。\n「病気の解説」…概要・感染経路・世界の状況・検査・予防・治療。\n「その他・資料の案内」…その他の感染症の簡単な紹介と、さらに詳しく調べるための資料。\n「見るときの注意」…数字を読むときに気をつけること。\n「出典一覧」…すべての資料名とURL。"),
    ("確定値", "国が1年分の届出を集計し直した数。"),
    ("速報値", "毎週や3か月ごとに出る途中の数。あとから変わることがあります。表では黄色で示しています。"),
    ("届出数", "医師が感染症と診断して保健所に届け出た数。検査を受けていない人などは入らないため、実際に感染している人の数とは同じではありません。"),
    ("罹患率（りかんりつ）", "人口10万人あたり、1年間に新しく病気になった人の数。人口の違う地域や年どうしを比べやすくなります。"),
    ("公費（こうひ）", "国や都道府県・市区町村が、税金などで費用を負担すること。"),
]
r = 3
for k, v in rows:
    put(intro, r, 1, k, f_bold, fill_soft, align=Alignment(vertical="top"))
    put(intro, r, 2, v, align=wrap)
    intro.row_dimensions[r].height = max(20, 16 * (v.count("\n") + math.ceil(len(v) / 64)))
    r += 1

wb.save(out_path)
print("saved", out_path)
