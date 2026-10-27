"""UNAIDSの国別のHIV推計（全部の国）を取ってきて、data/research/world_hiv_all.json を作る。

- 資料：UNAIDS「HIV estimates with uncertainty bounds 1990-2025」（2026年版の推計、2026-07-27公表）の表計算ファイル
- 使うのは2025年の、感染している人（推計の幅つき）・15〜49歳の感染率・新たな感染・エイズによる死亡と、
  感染を知っている人・治療を受けている人・ウイルスが抑えられている人の割合（どれもHIVとともに生きる人全体のうちの割合）
- 「2010年からの増減（%）」はこのファイルにない（UNAIDSは国別資料〔ファクトシート〕にだけ載せている）ため作らない
- 推計の数字がひとつもない国（日本・中国・アメリカなど）は入れない

使い方（感染症フォルダで）: python -I build/fetch_unaids_estimates.py （ダウンロードは build/_dl に置く。--refresh で取り直す）
"""
import json
import os
import re
import sys
import urllib.request

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DL = os.path.join(ROOT, "build", "_dl")
OUT = os.path.join(ROOT, "data", "research", "world_hiv_all.json")
CHECKED = "2026-10-27"
YEAR = 2025
PAGE = "https://www.unaids.org/en/resources/documents/2026/HIV_estimates_with_uncertainty_bounds_1990-present"
XLSX = "https://www.unaids.org/sites/default/files/2026-07/HIV_estimates_from_1990-to-present_1.xlsx"
os.makedirs(DL, exist_ok=True)

path = os.path.join(DL, "unaids_estimates_2026.xlsx")
if "--refresh" in sys.argv or not os.path.exists(path):
    req = urllib.request.Request(XLSX, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=180) as r, open(path, "wb") as f:
        f.write(r.read())
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)


def val(v):
    """「41.0m」「60 000」「<100」「>98」「...」などを数か文字にする。「<」「>」のついた値は文字のまま（UNAIDSの表し方）"""
    if v is None:
        return None
    s = str(v).replace(" ", " ").strip()
    if s in ("", "...", "-"):
        return None
    if s[0] in "<>":
        return s.replace(" ", "")
    s = s.replace(" ", "")
    m = re.fullmatch(r"([\d.]+)m", s)
    if m:
        return int(round(float(m.group(1)) * 1_000_000))
    return float(s) if "." in s else int(s)


# 推計の表：5行目が項目名、6行目が Estimate/Low/High。1列目が年、2列目がコード、3列目が名前
rows_e = list(wb["HIV2026Estimates_ByYear"].iter_rows(values_only=True))
e_head = [str(x).strip() if x else None for x in rows_e[4]]
E = {(r[0], r[1]): r for r in rows_e[6:] if r[0]}
EC = {k: e_head.index(t) for k, t in {
    "plhiv": "Estimated adults and children living with HIV",
    "prev": "Adults (15-49) prevalence (%)",
    "new": "Adults and children newly infected with HIV",
    "deaths": "AIDS-related deaths among adults and children",
}.items()}

# 検査と治療の表：4行目が項目名、5行目が年齢などの区分、6行目が Estimate/Low/High。「すべての年齢」の列を使う
rows_t = list(wb["HIV-Test-&-Treat_ByYear"].iter_rows(values_only=True))
T = {(r[0], r[1]): r for r in rows_t[6:] if r[0]}


def t_col(title):
    top = grp = None
    for j, (a, b, c) in enumerate(zip(rows_t[3], rows_t[4], rows_t[5])):
        top = str(a).replace("\n", " ") if a else top
        grp = str(b).strip() if b else grp
        if c == "Estimate" and top and top.startswith(title) and grp == "(All ages)":
            return j
    raise KeyError(title)


TC = {k: t_col(t) for k, t in {
    "know": "Percentage of people living with HIV who know their HIV status",
    "art": "Percentage of people living with HIV who are on antiretroviral treatment",
    "vls": "Percentage of people living with HIV with suppressed viral load",
}.items()}

out_rows = []
for code in sorted({c for (y, c) in E if y == YEAR}):
    if not re.fullmatch(r"[A-Z]{3}", code):
        continue  # 世界・地域（03M49WLD・UNAAP など）は入れない
    e, t = E[(YEAR, code)], T.get((YEAR, code)) or [None] * len(rows_t[5])
    row = {"iso3": code, "en": e[2], **{k: val(e[j]) for k, j in EC.items()}}
    row["plhiv_low"], row["plhiv_high"] = val(e[EC["plhiv"] + 1]), val(e[EC["plhiv"] + 2])
    row.update({k: val(t[j]) for k, j in TC.items()})
    if all(row[k] is None for k in [*EC, *TC]):
        continue
    out_rows.append(row)

out = {
    "checked": CHECKED,
    "data_year": YEAR,
    "rows": out_rows,
    "source": {"id": "S15", "title": f"HIV estimates with uncertainty bounds 1990-{YEAR}（全部の国の{YEAR}年の推計）", "publisher": "国連合同エイズ計画（UNAIDS）",
               "url": PAGE, "published": "2026-07-27（2026年版の推計）", "accessed": CHECKED, "data": XLSX},
}
json.dump(out, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("国の数", len(out_rows), "→", OUT)
